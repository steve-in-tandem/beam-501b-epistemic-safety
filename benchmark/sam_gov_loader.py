#!/usr/bin/env python3
"""
Beam-501B Epistemic Safety & RLVR "Proof of Truth" Benchmark Loader
Module: benchmark/sam_gov_loader.py

Downloads full-stack SAM.gov solicitation packets (opportunity metadata, notice
description, and ALL attached solicitation documents: SOWs, RFPs, Amendments,
Wage Determinations, Section L/M instructions, Q&A matrices, and Specs) for 10
real federal solicitations (ranging from 3 to 14 attachments each).

Extracts full text from every PDF, DOCX, XLSX, CSV, and TXT attachment, computes
cryptographic SHA-256 digests for every file and extracted text block, and writes
per-solicitation manifests to `benchmark/sam_packets/<notice_id>/manifest.json`.
"""

from __future__ import annotations

import csv
import hashlib
import html as htmllib
import io
import json
import os
import re
from pathlib import Path
from typing import Any, Dict, List

import docx
from docx.table import Table
import httpx
import openpyxl
from pypdf import PdfReader


ROOT_DIR = Path(__file__).resolve().parent.parent
ENV_PATH = ROOT_DIR / ".env"
SAM_PACKETS_DIR = ROOT_DIR / "benchmark" / "sam_packets"

# 10 real, active multi-attachment SAM.gov solicitations (3 to 14 attachments each)
TARGET_NOTICE_IDS: List[str] = [
    "f8fec0f1510d4ba68481131718e70063",  # 36C26026Q0827 (VA Puget Sound Grout Maintenance - 3 attachments)
    "fe284223457e4b30becd023a18ee036b",  # FA811827RB001 (USAF B-1 Electro-Mechanical Actuators - 3 attachments)
    "fdc18e6d405e49cc900ad85dbf967792",  # W912BU26QA035 (USACE Responder Uniform 5-Year IDC - 7 attachments)
    "fa6be0be3b904601989b8b64adec5afd",  # W912GY27QA002 (US Army Thermal Heat Boards - 4 attachments)
    "f5e5151c1caa4b5eb5ea7aeb814dbb9d",  # SPRTA127Q0033 (DLA Turbine Nozzle Support - 6 attachments)
    "f52092984da34f32936096d7a6d29e88",  # N4008527R0009 (US Navy NAVFAC Pest Control Services - 11 attachments)
    "f47c0a4ee61947e7a0edfbde89968cf8",  # 1131PL27R0001 (USTDA Reverse Trade Missions IDIQ - 5 attachments)
    "f4008c25a6024723bec93d2de70d059e",  # RFQHT001426QE048 (Defense Health Agency Hospital Comms - 11 attachments)
    "f281285f0050472aa3c8e507c7152076",  # 70Z03827QB0000011 (US Coast Guard MH-65 Trans Pump Cap - 3 attachments)
    "f7897da70a194ce4a422f679b585bd5e",  # 36C26126Q1261 (VA FY27 NRM Re-slurry & Re-stripe - 14 attachments)
]


def load_dotenv() -> None:
    if not ENV_PATH.exists():
        return
    for line in ENV_PATH.read_text(encoding="utf-8").splitlines():
        s = line.strip()
        if not s or s.startswith("#") or "=" not in s:
            continue
        k, v = s.split("=", 1)
        k = k.strip()
        v = v.strip().strip("'").strip('"')
        if k and v and k not in os.environ:
            os.environ[k] = v


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_text(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def safe_filename(name: str) -> str:
    name = (name or "attachment").replace("\\", "/").split("/")[-1]
    name = re.sub(r"[^\w.\- ()]+", "_", name).strip(" .") or "attachment"
    return name[:180]


def html_to_text(raw: str) -> str:
    if not raw:
        return ""
    t = re.sub(r"<li\b[^>]*>", "\n- ", raw, flags=re.I)
    t = re.sub(r"<(?:br|/p|/div|/li|/h[1-6]|/tr)\b[^>]*>", "\n", t, flags=re.I)
    t = re.sub(r"<[^>]{1,300}>", "", t)
    t = htmllib.unescape(t)
    t = re.sub(r"[ \t]+", " ", t)
    return re.sub(r"\n\s*\n\s*\n+", "\n\n", t).strip()


def extract_docx_text(data: bytes) -> str:
    d = docx.Document(io.BytesIO(data))
    parts: List[str] = []
    seen_hf = set()
    for section in d.sections:
        for hf in (section.header, section.footer):
            txt = " ".join(p.text for p in hf.paragraphs).strip()
            if txt and txt not in seen_hf:
                seen_hf.add(txt)
                parts.append(f"[Header/Footer] {txt}")
    for block in d.iter_inner_content():
        if isinstance(block, Table):
            for row in block.rows:
                cells: List[str] = []
                for c in row.cells:
                    txt = " ".join(c.text.split())
                    if not cells or cells[-1] != txt:
                        cells.append(txt)
                if any(cells):
                    parts.append(" | ".join(cells))
        elif block.text.strip():
            parts.append(block.text.strip())
    return "\n".join(parts)


def extract_xlsx_text(data: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out: List[str] = []
    try:
        for sheet in wb.worksheets:
            out.append(f"## Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if v is None else " ".join(str(v).split()) for v in row]
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    out.append(" | ".join(cells))
    finally:
        wb.close()
    return "\n".join(out)


def extract_pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        reader.decrypt("")
    pages: List[str] = []
    for idx, page in enumerate(reader.pages):
        txt = (page.extract_text() or "").strip()
        if txt:
            pages.append(f"--- [Page {idx + 1}] ---\n{txt}")
    return "\n\n".join(pages)


def extract_attachment_text(filename: str, data: bytes) -> str:
    lower = filename.lower()
    try:
        if lower.endswith(".pdf"):
            return extract_pdf_text(data)
        if lower.endswith(".docx"):
            return extract_docx_text(data)
        if lower.endswith(".xlsx") or lower.endswith(".xlsm"):
            return extract_xlsx_text(data)
        if lower.endswith(".csv"):
            rows = csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace")))
            return "\n".join(" | ".join(c.strip() for c in r) for r in rows if any(c.strip() for c in r))
        if lower.endswith(".txt") or lower.endswith(".md"):
            return data.decode("utf-8", errors="replace")
    except Exception as exc:
        return f"[Text extraction warning for {filename}: {exc}]"
    return ""


def fetch_solicitation_packet(client: httpx.Client, notice_id: str, sam_api_key: str) -> Dict[str, Any]:
    packet_dir = SAM_PACKETS_DIR / notice_id
    raw_dir = packet_dir / "attachments"
    raw_dir.mkdir(parents=True, exist_ok=True)

    # 1. Fetch opportunity metadata from SAM.gov v2 search API (using SAM_GOV_API_KEY) + public detail endpoint
    search_meta: Dict[str, Any] = {}
    if sam_api_key:
        r_search = client.get(
            "https://api.sam.gov/opportunities/v2/search",
            params={"api_key": sam_api_key, "noticeid": notice_id, "limit": 1},
        )
        if r_search.status_code == 200:
            opps = r_search.json().get("opportunitiesData") or []
            if opps:
                search_meta = opps[0]

    r_detail = client.get(f"https://sam.gov/api/prod/opps/v2/opportunities/{notice_id}")
    detail_json = r_detail.json() if r_detail.status_code == 200 else {}
    data2 = detail_json.get("data2") or {}

    title = search_meta.get("title") or data2.get("title") or "Untitled Solicitation"
    sol_num = (
        search_meta.get("solicitationNumber")
        or data2.get("solicitationNumber")
        or notice_id[:12].upper()
    )
    agency = (
        search_meta.get("fullParentPathName")
        or search_meta.get("department")
        or "Federal Agency"
    )
    naics_code = search_meta.get("naicsCode") or (
        (data2.get("naics") or [{}])[0].get("code", [""])[0]
        if data2.get("naics")
        else ""
    )
    posted_date = search_meta.get("postedDate") or detail_json.get("postedDate") or ""

    # Extract HTML description from detail_json['description'] list
    desc_blocks = detail_json.get("description") or []
    raw_desc_html = "\n\n".join(b.get("body", "") for b in desc_blocks if isinstance(b, dict))
    description_text = html_to_text(raw_desc_html)

    # 2. List all attachments via SAM's v3 resources endpoint
    r_res = client.get(
        f"https://sam.gov/api/prod/opps/v3/opportunities/{notice_id}/resources"
    )
    res_json = r_res.json() if r_res.status_code == 200 else {}
    attachment_lists = (res_json.get("_embedded") or {}).get(
        "opportunityAttachmentList"
    ) or []

    attachments_manifest: List[Dict[str, Any]] = []
    for lst in attachment_lists:
        for att in lst.get("attachments") or []:
            if att.get("deletedFlag") == "1" or att.get("type") != "file":
                continue
            if (att.get("accessLevel") or "public") != "public":
                continue

            resource_id = att.get("resourceId")
            orig_name = att.get("name") or f"{resource_id}.bin"
            clean_name = safe_filename(orig_name)
            download_url = f"https://sam.gov/api/prod/opps/v3/opportunities/resources/files/{resource_id}/download"

            local_file_path = raw_dir / clean_name
            if local_file_path.exists() and local_file_path.stat().st_size > 0:
                file_bytes = local_file_path.read_bytes()
            else:
                dl_resp = client.get(download_url, follow_redirects=True)
                if dl_resp.status_code != 200:
                    print(f"    [WARN] Failed to download {clean_name} (HTTP {dl_resp.status_code})")
                    continue
                file_bytes = dl_resp.content
                local_file_path.write_bytes(file_bytes)

            extracted_text = extract_attachment_text(clean_name, file_bytes)
            attachments_manifest.append(
                {
                    "resource_id": resource_id,
                    "filename": clean_name,
                    "size_bytes": len(file_bytes),
                    "file_sha256": sha256_bytes(file_bytes),
                    "extracted_chars": len(extracted_text),
                    "extracted_sha256": sha256_text(extracted_text),
                    "extracted_text": extracted_text,
                }
            )
            print(
                f"    + Downloaded {clean_name} ({len(file_bytes):,} bytes -> {len(extracted_text):,} chars extracted)"
            )

    manifest = {
        "notice_id": notice_id,
        "sam_url": f"https://sam.gov/workspace/contract/opp/{notice_id}/view",
        "solicitation_number": str(sol_num).strip(),
        "title": str(title).strip(),
        "agency": str(agency).strip(),
        "naics_code": str(naics_code).strip(),
        "posted_date": str(posted_date).strip(),
        "description_text": description_text,
        "attachment_count": len(attachments_manifest),
        "attachments": attachments_manifest,
    }

    manifest_path = packet_dir / "manifest.json"
    manifest_path.write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    return manifest


def main() -> None:
    load_dotenv()
    sam_api_key = os.environ.get("SAM_GOV_API_KEY", "")
    SAM_PACKETS_DIR.mkdir(parents=True, exist_ok=True)

    print("==============================================================================")
    print("SAM.gov Full-Stack Multi-Attachment Solicitation Downloader (N=10)")
    print(f"SAM_GOV_API_KEY configured: {'YES' if sam_api_key else 'NO'}")
    print("==============================================================================")

    summaries: List[Dict[str, Any]] = []
    with httpx.Client(
        timeout=httpx.Timeout(90.0, connect=20.0),
        headers={"User-Agent": "InTandem-Beam501B-SafetyHarness/1.0"},
    ) as client:
        for idx, nid in enumerate(TARGET_NOTICE_IDS, start=1):
            print(f"\n[{idx}/10] Fetching Notice {nid} ...")
            manifest = fetch_solicitation_packet(client, nid, sam_api_key)
            total_chars = len(manifest["description_text"]) + sum(
                a["extracted_chars"] for a in manifest["attachments"]
            )
            print(
                f"  => {manifest['solicitation_number']} | {manifest['title'][:58]} | {manifest['attachment_count']} attachments | {total_chars:,} total chars"
            )
            summaries.append(
                {
                    "notice_id": nid,
                    "solicitation_number": manifest["solicitation_number"],
                    "title": manifest["title"],
                    "agency": manifest["agency"],
                    "attachment_count": manifest["attachment_count"],
                    "total_extracted_chars": total_chars,
                }
            )

    index_path = SAM_PACKETS_DIR / "index.json"
    index_path.write_text(json.dumps(summaries, indent=2) + "\n", encoding="utf-8")
    print("\n==============================================================================")
    print(f"Successfully downloaded and extracted {len(summaries)} full-stack SAM.gov packets!")
    print(f"Total attachments downloaded & extracted: {sum(s['attachment_count'] for s in summaries)}")
    print("==============================================================================")


if __name__ == "__main__":
    main()

